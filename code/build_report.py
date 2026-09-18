# -*- coding: utf-8 -*-
"""
竞赛论文报告生成脚本：
  1) 绘制技术路线图 fig0
  2) 从 results/计算结果汇总.xlsx 读取全部真实计算结果
  3) 自动排版生成符合数模论文体例的 Word：封面、摘要、目录、正文、参考文献、附录
运行：python build_report.py  （需先运行 main.py 生成结果与图形）
"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(BASE, ".."))
FIG = os.path.join(ROOT, "figures")
RES = os.path.join(ROOT, "results")
XLSX = os.path.join(RES, "计算结果汇总.xlsx")
OUT = os.path.join(ROOT, "我国省域数字经济发展水平的综合评价、障碍诊断与趋势预测.docx")

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

# =============================================================================
# 0. 绘制技术路线图 fig0
# =============================================================================
def draw_flowchart():
    fig, ax = plt.subplots(figsize=(10.5, 7.6))
    ax.set_xlim(0, 12); ax.set_ylim(0, 12); ax.axis("off")

    def box(x, y, w, h, text, fc, fs=12, tc="black"):
        p = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.08,rounding_size=0.12",
                           linewidth=1.4, edgecolor="#34495e", facecolor=fc)
        ax.add_patch(p)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
                fontsize=fs, color=tc, wrap=True)

    def arrow(x1, y1, x2, y2):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                     mutation_scale=16, lw=1.4, color="#34495e"))

    box(3.0, 10.6, 6.0, 1.0, "数据层：31省数字经济指标数据采集\n（《中国统计年鉴》、工信部、数字普惠金融指数）",
        "#d6eaf8", 11)
    box(3.0, 8.9, 6.0, 1.0, "数据预处理：指标正向化 → 极差标准化 → 构建评价矩阵", "#d6eaf8", 11)
    arrow(6, 10.6, 6, 9.9)

    # 三个问题并列
    box(0.3, 6.4, 3.6, 1.7, "问题一  综合评价\n熵权法赋权\nTOPSIS 贴近度排序\nWard 系统聚类分梯队",
        "#fdebd0", 11)
    box(4.2, 6.4, 3.6, 1.7, "问题二  障碍诊断\n构建障碍度模型\n指标层 / 维度层\n主要制约因子识别",
        "#fdebd0", 11)
    box(8.1, 6.4, 3.6, 1.7, "问题三  趋势预测\n灰色 GM(1,1) 模型\n后验差比 C / 小误差概率 P\n预测 2024—2028",
        "#fdebd0", 11)
    for x in (2.1, 6.0, 9.9):
        arrow(6, 8.9, x, 8.1)

    box(1.4, 4.2, 9.2, 1.1, "结果层：发展水平排序与三大梯队格局 · 关键障碍因子 · 未来五年演化趋势",
        "#d5f5e3", 11.5)
    for x in (2.1, 6.0, 9.9):
        arrow(x, 6.4, 6, 5.3)

    box(3.0, 2.2, 6.0, 1.1, "问题四  分梯队差异化政策建议\n（巩固引领区 · 攻坚追赶区 · 补齐欠发达区）",
        "#fadbd8", 11.5)
    arrow(6, 4.2, 6, 3.3)
    box(3.0, 0.4, 6.0, 1.0, "目标：为缩小区域数字鸿沟、推动数字经济协调发展提供决策依据",
        "#eaecee", 11)
    arrow(6, 2.2, 6, 1.4)
    fig.tight_layout()
    fp = os.path.join(FIG, "fig0_技术路线图.png")
    fig.savefig(fp, dpi=300, bbox_inches="tight"); plt.close(fig)
    return fp


# =============================================================================
# 1. 样式与基础排版工具
# =============================================================================
CN_SONG, CN_HEI, EN_FONT = "宋体", "黑体", "Times New Roman"

def set_font(run, cn=CN_SONG, en=EN_FONT, size=12, bold=False, color=(0, 0, 0), italic=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:ascii"), en); rf.set(qn("w:hAnsi"), en)

def init_styles(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT; normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_SONG)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    for lvl, sz in [(1, 16), (2, 14), (3, 12)]:
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Arial"; st.font.size = Pt(sz); st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEI)

def h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), cn=CN_HEI, en="Arial", size=16, bold=True)
    return p

def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(11); p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), cn=CN_HEI, en="Arial", size=14, bold=True)
    return p

def body(doc, text, indent=True):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    set_font(p.add_run(text), size=12)
    return p

def formula(doc, text, num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text + "          （" + num + "）"), en="Cambria Math", size=12)
    return p

def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), cn=CN_SONG, size=10.5)
    return p

def add_figure(doc, fname, cap, width=14.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(FIG, fname), width=Cm(width))
    caption(doc, cap)

def _set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)

def add_table(doc, df, fontsize=10.5, header_fill="F2F2F2", align_center_cols=None,
              col_widths=None, bold_last_row=False):
    """df 列名与内容直接写入；首行为表头并在跨页时重复。"""
    nrow, ncol = df.shape[0] + 1, df.shape[1]
    tb = doc.add_table(rows=nrow, cols=ncol)
    tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = True
    align_center_cols = align_center_cols or set()
    # 表头
    _repeat_header(tb.rows[0])
    for j, col in enumerate(df.columns):
        c = tb.cell(0, j); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_bg(c, header_fill)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        set_font(p.add_run(str(col)), cn=CN_HEI, size=fontsize, bold=True)
    # 数据
    for i in range(df.shape[0]):
        for j in range(ncol):
            c = tb.cell(i + 1, j); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (j == 0 or j in align_center_cols) \
                else WD_ALIGN_PARAGRAPH.CENTER
            val = df.iat[i, j]
            set_font(p.add_run(str(val)), size=fontsize,
                     bold=(bold_last_row and i == df.shape[0] - 1))
            if i % 2 == 1:
                _set_cell_bg(c, "FAFAFA")
    # 禁止同一行跨页断开
    for r in tb.rows:
        trPr = r._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    if col_widths:
        for j, w in enumerate(col_widths):
            for r in tb.rows:
                r.cells[j].width = Cm(w)
    return tb

def add_field(paragraph, field_code):
    """插入 Word 域（如 PAGE / TOC）。"""
    run = paragraph.add_run()
    fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); it.text = field_code
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = ""
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r = run._element
    r.append(fb); r.append(it); r.append(sep); r.append(txt); r.append(end)
    return run

def enable_update_fields(doc):
    s = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); s.append(uf)

from docx.enum.text import WD_BREAK
def pagebreak(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)


# =============================================================================
# 2. 读取计算结果
# =============================================================================
def load_results():
    xl = pd.ExcelFile(XLSX)
    R = {name: pd.read_excel(xl, name) for name in xl.sheet_names}
    return R


# =============================================================================
# 3. 生成文档
# =============================================================================
def build():
    fp0 = draw_flowchart()
    R = load_results()
    doc = Document()
    init_styles(doc)
    enable_update_fields(doc)

    # ---------------------------------------------------------------- 封面
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("校大学生数学建模竞赛"), cn=CN_HEI, en="Arial", size=22, bold=True)
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("我国省域数字经济发展水平的"), cn=CN_HEI, en="Arial", size=20, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("综合评价、障碍诊断与趋势预测"), cn=CN_HEI, en="Arial", size=20, bold=True)
    for _ in range(5):
        doc.add_paragraph()
    info = [("参赛编号：", "____________________"),
            ("参赛队员：", "____________  ____________  ____________"),
            ("指导教师：", "____________________"),
            ("完成日期：", "2026 年 9 月")]
    for k, v in info:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(k + v), size=14)
    # 封面所在第 1 节，不显示页码
    sec_cover = doc.sections[0]

    # ---------------------------------------------------------------- 摘要页（第 2 节）
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("摘    要"), cn=CN_HEI, en="Arial", size=16, bold=True)

    abs_paras = [
        "数字经济已成为推动我国经济高质量发展的核心引擎，但省域之间在信息基础设施、数字产业、"
        "融合应用与创新能力上发展并不均衡。本文围绕“评价—诊断—预测—对策”的逻辑主线，构建涵盖"
        "数字基础设施、数字产业化、产业数字化、数字创新能力和数字普惠五个维度、共 11 项指标的省域"
        "数字经济发展水平评价指标体系，综合运用熵权法、TOPSIS、系统聚类、障碍度模型和灰色 GM(1,1) "
        "模型，对我国 31 个省（区、市）展开定量研究。",
        "针对问题一，先对负向指标进行正向化、对全部指标做极差标准化，再用熵权法依据指标信息量客观"
        "赋权，避免主观偏差。结果表明软件业务收入（11.641%）、企业电子商务销售额（11.388%）和信息"
        "技术领域专利授权（10.992%）权重最高；维度上以数字基础设施（26.740%）和数字产业化（22.105%）"
        "贡献最大。进而用 TOPSIS 计算相对贴近度并排序，北京（0.958）、广东（0.941）、上海（0.936）、"
        "江苏（0.918）、浙江（0.872）位列前五；结合 Ward 系统聚类将 31 省划分为三个梯队：第Ⅰ梯队 5 省、"
        "第Ⅱ梯队 14 省、第Ⅲ梯队 12 省，整体呈现“东部引领、中部追赶、西部偏弱”的梯度格局。",
        "针对问题二，构建障碍度模型测算各指标对发展水平的制约程度。指标层面，软件业务收入（12.651%）、"
        "信息技术专利授权（12.436%）、电子商务销售额（11.720%）和 5G 基站数（10.328%）是全国最主要的"
        "障碍因子；维度层面，数字基础设施（26.834%）和数字产业化（22.049%）制约最强，说明“硬件底座”"
        "与“核心产业”仍是短板集中区，且不同梯队省份的首要障碍存在明显差异。",
        "针对问题三，对 2014—2023 年全国数字经济综合发展指数建立灰色 GM(1,1) 模型，求得发展系数 "
        "a=−0.11045、灰作用量 b=0.15936；后验差比值 C=0.0103、小误差概率 P=1.000、平均相对误差仅 "
        "0.217%，模型精度为一级（优）。预测显示 2024—2028 年指数将由 0.5084 稳步升至 0.7908，年均增长"
        "约 11.68%，保持强劲上升态势。",
        "针对问题四，依据评价与诊断结果，从巩固第Ⅰ梯队引领优势、推动第Ⅱ梯队错位攻坚、补齐第Ⅲ梯队"
        "数字基建与人才短板、强化国家层面顶层设计四个方面提出差异化、可操作的政策建议。本文模型客观、"
        "可复算、可推广，可为区域数字经济协调发展决策提供量化参考。",
    ]
    for t in abs_paras:
        body(doc, t)
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_before = Pt(8)
    set_font(p.add_run("关键词："), cn=CN_HEI, size=12, bold=True)
    set_font(p.add_run("数字经济；熵权法；TOPSIS；障碍度模型；灰色 GM(1,1)；综合评价"), size=12)

    # ---------------------------------------------------------------- 目录（第 3 节前为目录页）
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("目    录"), cn=CN_HEI, en="Arial", size=16, bold=True)
    toc_p = doc.add_paragraph()
    add_field(toc_p, r' TOC \o "1-2" \h \z \u ')

    # ---------------------------------------------------------------- 正文（第 3 节，页码从 1 开始）
    body_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    # 断开与前节页脚链接，正文页脚居中插入 PAGE 域并重新编号
    body_sec.footer.is_linked_to_previous = False
    fp = body_sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(add_field(fp, " PAGE "), size=10.5)
    sectPr = body_sec._sectPr
    pn = OxmlElement("w:pgNumType"); pn.set(qn("w:start"), "1"); sectPr.append(pn)
    # 前两节页脚置空
    for s in doc.sections[:2]:
        s.footer.is_linked_to_previous = False

    # 一、问题重述
    h1(doc, "一、问题重述")
    h2(doc, "（一）问题背景")
    body(doc, "当前，以数据为关键生产要素、以现代信息网络为重要载体的数字经济正深刻重塑生产方式、"
              "生活方式和治理方式。国家“十四五”规划和《数字中国建设整体布局规划》明确提出要做强做优"
              "做大数字经济。然而，由于资源禀赋、产业基础和区位条件不同，我国各省数字经济发展水平存在"
              "显著的空间差异。科学测度省域数字经济发展水平、精准识别制约因素、合理预判未来趋势，对于"
              "实施差异化区域政策、缩小区域数字鸿沟具有重要的现实意义。")
    body(doc, "本题给出我国 31 个省（自治区、直辖市，不含港澳台）数字基础设施、数字产业、融合应用、"
              "创新能力和数字普惠等方面的指标数据，以及 2014—2023 年全国数字经济综合发展指数时间序列，"
              "要求建立数学模型完成评价、诊断、预测并提出建议。")
    h2(doc, "（二）问题提出")
    for t in [
        "问题一：构建省域数字经济发展水平评价指标体系，确定各指标权重，对 31 个省份进行综合评价与"
        "排序，并划分发展梯队、刻画空间格局。",
        "问题二：在综合评价基础上，识别制约各省及全国数字经济发展的主要障碍因子，区分指标层与维度层"
        "的障碍程度。",
        "问题三：依据 2014—2023 年全国数字经济综合发展指数，建立预测模型对 2024—2028 年总体发展趋势"
        "进行外推预测，并检验模型精度。",
        "问题四：综合上述定量结果，针对不同发展梯队提出差异化、可操作的政策建议。"]:
        body(doc, t)

    # 二、问题分析
    h1(doc, "二、问题分析")
    h2(doc, "（一）总体思路与技术路线")
    body(doc, "本文遵循“数据预处理—综合评价—障碍诊断—趋势预测—对策建议”的逻辑主线。首先对原始指标"
              "进行正向化与无量纲化，消除量纲与方向差异；问题一属于多指标综合评价问题，采用客观赋权的"
              "熵权法确定权重，结合 TOPSIS 计算相对贴近度排序，并以 Ward 系统聚类划分梯队；问题二属于"
              "诊断型问题，引入障碍度模型量化各因子的制约强度；问题三属于小样本、贫信息的时间序列预测"
              "问题，灰色 GM(1,1) 模型对短序列、单调趋势数据适应性强，故采用之并做严格精度检验；问题四"
              "在前述结果上分梯队施策。总体技术路线如图 1 所示。")
    add_figure(doc, "fig0_技术路线图.png", "图1  本文研究技术路线图", width=14.0)
    h2(doc, "（二）各问题具体分析")
    body(doc, "对于问题一，评价的关键在于权重确定与综合排序。层次分析法等主观赋权易受专家偏好影响，"
              "而熵权法完全依据指标数据的离散程度（信息量）赋权，离散程度越大、权重越高，客观性强；"
              "TOPSIS 通过度量各评价对象与正、负理想解的距离计算贴近度，对样本量无严格要求、结果直观，"
              "二者结合可实现“客观赋权 + 优劣排序”。在此基础上用 Ward 法对贴近度系统聚类，可避免人为"
              "划定梯队阈值的随意性。")
    body(doc, "对于问题二，障碍度模型能够把“权重（重要性）”与“指标差距（短板程度）”结合起来，既考虑"
              "因子的重要程度，又考虑其当前与理想状态的偏离，从而准确识别真正“卡脖子”的环节，并可在"
              "指标层和维度层两个层级上分别汇总。")
    body(doc, "对于问题三，全国数字经济综合发展指数仅有 10 个年度观测，样本量小、且近似单调指数增长，"
              "正契合灰色系统理论“小样本、贫信息”建模的适用条件。GM(1,1) 通过一次累加弱化随机性、"
              "拟合指数规律，再累减还原，适合中短期外推；同时用后验差比值 C、小误差概率 P 和平均相对"
              "误差检验其可靠性。")

    # 三、模型假设
    h1(doc, "三、模型假设")
    ass = [
        "假设所采用的官方统计数据真实可靠，统计口径在研究期内保持一致，不存在系统性错报或漏报；",
        "假设研究期内国家宏观政策与外部环境不发生颠覆性突变，数字经济发展保持总体连续、稳定的演化趋势；",
        "假设所选取的 11 项指标能够较为全面地刻画省域数字经济发展水平，指标之间不存在严重的多重共线性；",
        "假设各指标对综合发展水平均为单调影响，即正向指标越大越好、负向指标越小越好，不存在最优区间；",
        "假设灰色 GM(1,1) 模型的中短期外推成立，预测期内发展系数保持稳定。"]
    for i, t in enumerate(ass, 1):
        body(doc, f"{i}. {t}")

    # 四、符号说明
    h1(doc, "四、符号说明")
    sym = pd.DataFrame({
        "符号": ["x_ij / z_ij", "m_j / M_j", "p_ij", "e_j / g_j", "w_j",
                "V+ / V−", "D_i+ / D_i−", "C_i", "I_ij / O_ij",
                "x^(0) / x^(1)", "a / b", "S1 / S2", "C / P"],
        "含义": [
            "第 i 省第 j 项指标的原始值 / 标准化值",
            "第 j 项指标的最小值 / 最大值",
            "第 j 项指标下第 i 省的特征比重",
            "第 j 项指标的信息熵 / 差异系数",
            "第 j 项指标的熵权权重",
            "正理想解 / 负理想解",
            "第 i 省到正、负理想解的欧氏距离",
            "第 i 省的 TOPSIS 相对贴近度",
            "指标偏离度 / 障碍度",
            "原始序列 / 一次累加生成序列（1-AGO）",
            "GM(1,1) 发展系数 / 灰作用量",
            "原始序列标准差 / 残差序列标准差",
            "后验差比值 / 小误差概率"]})
    caption(doc, "表1  主要符号说明")
    add_table(doc, sym, fontsize=10.5, col_widths=[3.6, 12.0])

    # 五、数据来源与预处理
    h1(doc, "五、数据来源与预处理")
    h2(doc, "（一）数据来源与指标体系构建")
    body(doc, "本文数据主要来源于《中国统计年鉴》《中国信息产业年鉴》、工业和信息化部行业统计公报以及"
              "北京大学数字普惠金融指数。遵循科学性、系统性、可比性和数据可获得性原则，从数字基础设施、"
              "数字产业化、产业数字化、数字创新能力和数字普惠五个维度选取 11 项二级指标，构建“目标层—"
              "维度层—指标层”三级评价指标体系，其中除城乡数字接入鸿沟指数为负向指标外，其余均为正向"
              "指标，具体见表2。")
    idx_sys = pd.DataFrame({
        "维度层": ["数字基础设施", "数字基础设施", "数字基础设施", "数字产业化", "数字产业化",
                 "产业数字化", "产业数字化", "数字创新能力", "数字创新能力", "数字普惠", "数字普惠"],
        "编号": [f"X{i}" for i in range(1, 12)],
        "指标": ["每百人移动电话拥有量", "单位国土面积光缆长度", "5G基站数", "软件业务收入",
               "电信业务总量", "企业电子商务销售额", "关键工序数控化率", "R&D经费投入强度",
               "信息技术领域专利授权", "数字普惠金融指数", "城乡数字接入鸿沟指数"],
        "方向": ["正向"] * 10 + ["负向"]})
    caption(doc, "表2  省域数字经济发展水平评价指标体系")
    add_table(doc, idx_sys, fontsize=10.5, col_widths=[3.4, 1.8, 7.2, 2.2])
    h2(doc, "（二）指标正向化与无量纲化")
    body(doc, "由于各指标量纲不同、数量级差异悬殊，且存在负向指标，需先进行正向化与极差标准化，将全部"
              "数据映射到 [0,1] 区间。正向、负向指标分别按式（1）处理：")
    formula(doc, "正向：z_ij = (x_ij − m_j)/(M_j − m_j)；  负向：z_ij = (M_j − x_ij)/(M_j − m_j)", "1")
    body(doc, "式中 m_j、M_j 分别为第 j 项指标在 31 省中的最小值与最大值。标准化后各指标最大值为 1、"
              "最小值为 0，既消除了量纲，又保证后续熵权与距离计算的可比性。")

    # 六、问题一
    h1(doc, "六、问题一：熵权—TOPSIS 综合评价与梯队划分")
    h2(doc, "（一）熵权法确定权重")
    body(doc, "熵权法依据信息熵反映指标的离散程度：指标数据差异越大，信息熵越小、所包含的信息量越大，"
              "应赋予越高权重。计算步骤为：先由标准化矩阵计算特征比重 p_ij，再求信息熵 e_j、差异系数 "
              "g_j，最后归一化得到权重 w_j，如式（2）—（4）：")
    formula(doc, "p_ij = z_ij / Σᵢ z_ij", "2")
    formula(doc, "e_j = −(1/ln n) · Σᵢ p_ij · ln p_ij，  g_j = 1 − e_j", "3")
    formula(doc, "w_j = g_j / Σⱼ g_j，  且 Σⱼ w_j = 1", "4")
    body(doc, "其中 n=31 为省份数，并约定当 p_ij=0 时 p_ij·ln p_ij=0。依此计算 11 项指标的信息熵、差异"
              "系数与权重，结果见表3，五个维度的权重合计见表4。")
    wdf = R["熵权权重"][["指标编号", "指标名称", "所属维度", "信息熵e", "差异系数g", "权重w(%)"]].copy()
    caption(doc, "表3  各指标信息熵、差异系数与熵权权重")
    add_table(doc, wdf, fontsize=10, col_widths=[2.0, 4.6, 3.2, 2.0, 2.2, 2.0])
    dimw = R["维度权重"].copy(); dimw.columns = ["维度", "维度权重(%)"]
    caption(doc, "表4  五个一级维度权重合计")
    add_table(doc, dimw, fontsize=10.5, col_widths=[6.0, 6.0])
    body(doc, "由表3、表4可见：第一，软件业务收入、企业电子商务销售额、信息技术专利授权的权重位居前三，"
              "分别为 11.641%、11.388%、10.992%，说明数字核心产业规模与融合应用、技术创新产出是拉开"
              "省域差距的关键；第二，数字普惠金融指数权重最低（5.685%），表明经过多年推广，该指标在"
              "省际已较为均衡、区分度较小；第三，从维度看，数字基础设施（26.740%）与数字产业化（22.105%）"
              "合计接近一半权重，是数字经济竞争的主战场。")
    add_figure(doc, "fig1_熵权权重.png", "图2  数字经济发展水平评价指标熵权权重", width=13.5)

    h2(doc, "（二）TOPSIS 综合评价模型")
    body(doc, "TOPSIS（逼近理想解排序法）通过构造加权标准化矩阵，确定正理想解 V⁺（各指标最优值）与负"
              "理想解 V⁻（各指标最劣值），分别计算各省到二者的欧氏距离 D_i⁺、D_i⁻，进而得到相对贴近度 "
              "C_i，C_i 越接近 1 表明发展水平越高，计算见式（5）—（7）：")
    formula(doc, "v_ij = w_j · z_ij，  V⁺ = {max v_ij}，  V⁻ = {min v_ij}", "5")
    formula(doc, "D_i⁺ = √(Σⱼ(v_ij − v_j⁺)²)，  D_i⁻ = √(Σⱼ(v_ij − v_j⁻)²)", "6")
    formula(doc, "C_i = D_i⁻ /(D_i⁺ + D_i⁻)，  C_i ∈ [0,1]", "7")
    body(doc, "31 省贴近度计算结果及排序见表5，可视化结果见图3。")
    rank = R["TOPSIS评价结果"][["排名", "省份", "正理想距离D+", "负理想距离D-", "贴近度C", "梯队"]].copy()
    caption(doc, "表5  31 省 TOPSIS 综合评价结果（按贴近度降序）")
    add_table(doc, rank, fontsize=9.5,
              col_widths=[1.6, 2.6, 3.0, 3.0, 2.4, 2.6])
    add_figure(doc, "fig2_TOPSIS贴近度排名.png", "图3  31 省数字经济发展水平 TOPSIS 贴近度及梯队划分", width=13.0)

    h2(doc, "（三）Ward 系统聚类与梯队划分")
    body(doc, "为避免主观设定梯队分界阈值，本文以相对贴近度为变量，采用 Ward 最小方差法进行系统聚类，"
              "并依据类内平均贴近度由高到低将 31 省划分为第Ⅰ、Ⅱ、Ⅲ梯队，结果见表6、图4。")
    tier_tbl = pd.DataFrame({
        "梯队": ["第Ⅰ梯队", "第Ⅱ梯队", "第Ⅲ梯队"],
        "省份数量": [5, 14, 12],
        "占比(%)": [16.1, 45.2, 38.7],
        "平均贴近度": [0.925, 0.437, 0.134],
        "包含省份": [
            "北京、广东、上海、江苏、浙江",
            "福建、山东、天津、重庆、湖北、四川、安徽、陕西、湖南、河南、辽宁、江西、河北、海南",
            "山西、内蒙古、黑龙江、吉林、广西、宁夏、新疆、贵州、云南、甘肃、青海、西藏"]})
    caption(doc, "表6  三大发展梯队划分结果")
    add_table(doc, tier_tbl, fontsize=10, col_widths=[2.2, 1.8, 1.8, 2.2, 7.6])
    add_figure(doc, "fig6_梯队构成.png", "图4  三大发展梯队省份数量构成", width=9.5)
    body(doc, "结果呈现出鲜明的“金字塔型”梯度结构：第Ⅰ梯队为北京、广东、上海、江苏、浙江五省市，平均"
              "贴近度高达 0.925，数字经济全面领先；第Ⅱ梯队以中西部省会经济强省和东部沿海部分省份为主，"
              "平均贴近度 0.437，是承上启下的追赶群体；第Ⅲ梯队 12 省多位于西部和东北地区，平均贴近度仅"
              "0.134，西藏（0.000）、青海（0.058）、甘肃（0.067）处于末位。为对比不同梯队的内部结构，"
              "选取各梯队中贴近度最接近梯队均值的代表省份（江苏、安徽、新疆）绘制五维雷达图（图5），"
              "可见第Ⅰ梯队在五个维度上全面外扩，第Ⅱ梯队呈“数字普惠相对较好、核心产业偏弱”的特征，"
              "第Ⅲ梯队则各维度普遍收缩、短板交织。")
    add_figure(doc, "fig3_梯队代表雷达图.png", "图5  不同梯队代表省份五维度得分雷达图", width=11.0)

    # 七、问题二
    h1(doc, "七、问题二：基于障碍度模型的制约因素诊断")
    h2(doc, "（一）障碍度模型构建")
    body(doc, "障碍度模型通过“因子贡献度、指标偏离度、障碍度”三个变量刻画单项指标对总体发展水平的制约"
              "程度。因子贡献度取熵权 w_j（表示重要性），指标偏离度 I_ij=1−z_ij（表示当前值距最优值的"
              "差距），障碍度计算见式（8），并可在维度层对所属指标求和：")
    formula(doc, "I_ij = 1 − z_ij，  O_ij = w_j·I_ij / Σⱼ(w_j·I_ij) × 100%", "8")
    h2(doc, "（二）指标层障碍度分析")
    obs_ind = R["指标障碍度"].copy()
    caption(doc, "表7  指标层平均障碍度排序（全国）")
    add_table(doc, obs_ind, fontsize=10, col_widths=[2.4, 6.6, 4.0])
    add_figure(doc, "fig4_障碍因子排序.png", "图6  制约数字经济发展的主要障碍因子（全国平均）", width=13.5)
    body(doc, "由表7、图6可知，全国层面前四位障碍因子依次为软件业务收入（12.651%）、信息技术领域专利"
              "授权（12.436%）、企业电子商务销售额（11.720%）和 5G 基站数（10.328%），累计贡献近 47% 的"
              "障碍度。它们集中在数字产业化、产业数字化和新型基础设施领域，说明制约我国数字经济整体"
              "提档升级的，已不再是移动电话、普惠金融等普及型指标，而是高端软件产业、核心技术专利、"
              "数字化交易规模和 5G 网络深度覆盖等“高质量供给”。")
    h2(doc, "（三）维度层障碍度分析")
    obs_dim = R["维度障碍度"].copy(); obs_dim.columns = ["维度", "维度平均障碍度(%)"]
    caption(doc, "表8  维度层平均障碍度排序")
    add_table(doc, obs_dim, fontsize=10.5, col_widths=[6.0, 6.0])
    body(doc, "维度层障碍度由高到低依次为数字基础设施（26.834%）、数字产业化（22.049%）、数字创新能力"
              "（19.668%）、产业数字化（18.643%）、数字普惠（12.806%）。这与熵权结论相互印证：越是权重"
              "高、省际差异大的维度，越是短板集中、需要重点突破的领域；数字普惠障碍度最低，说明数字"
              "服务的“广度覆盖”已基本实现，下一阶段应转向“深度提升”。")
    h2(doc, "（四）典型省份障碍因子对比")
    body(doc, "进一步选取排名首位、居中和末位的北京、辽宁、西藏三省对比其前三位障碍因子：北京为软件"
              "业务收入（22.53%）、5G 基站数（21.41%）、单位面积光缆长度（14.75%），其短板表现为高密度"
              "地区网络承载与产业规模“再提升”的约束；辽宁为软件业务收入（12.61%）、电子商务销售额"
              "（12.46%）、信息技术专利（12.10%），障碍相对均衡、集中于产业与创新；西藏为软件业务收入"
              "（11.64%）、电子商务销售额（11.39%）、信息技术专利（10.99%），呈现“全域偏弱、产业尤甚”"
              "的特征。这表明不同发展阶段省份的矛盾重点不同，政策不能“一刀切”。")

    # 八、问题三
    h1(doc, "八、问题三：基于灰色 GM(1,1) 的趋势预测")
    h2(doc, "（一）GM(1,1) 模型原理")
    body(doc, "灰色 GM(1,1) 模型适用于小样本、单调变化序列的中短期预测。设原始非负序列为 x⁽⁰⁾，先做一次"
              "累加生成（1-AGO）x⁽¹⁾以弱化随机波动，并对 x⁽¹⁾ 作紧邻均值生成 z⁽¹⁾，建立白化微分方程，"
              "如式（9）—（11）：")
    formula(doc, "x⁽¹⁾(k) = Σₜ₌₁ᵏ x⁽⁰⁾(t)，  z⁽¹⁾(k) = 0.5[x⁽¹⁾(k) + x⁽¹⁾(k−1)]", "9")
    formula(doc, "x⁽⁰⁾(k) + a·z⁽¹⁾(k) = b，  [a,b]ᵀ = (BᵀB)⁻¹BᵀY", "10")
    formula(doc, "x̂⁽¹⁾(k+1) = [x⁽⁰⁾(1) − b/a]·e^(−ak) + b/a，  x̂⁽⁰⁾(k+1) = x̂⁽¹⁾(k+1) − x̂⁽¹⁾(k)", "11")
    body(doc, "其中 a 为发展系数（a<0 表示增长）、b 为灰作用量，求得时间响应函数后再累减还原即得拟合与"
              "预测值。")
    h2(doc, "（二）参数求解与精度检验")
    body(doc, "以 2014—2023 年全国数字经济综合发展指数为原始序列建模，最小二乘估计得发展系数 "
              "a=−0.11045、灰作用量 b=0.15936。采用残差检验、后验差检验和小误差概率检验综合评定精度，"
              "其中后验差比值 C=S₂/S₁（越小越好）、小误差概率 P（越大越好）。历史拟合结果见表9。")
    fit = R["GM11拟合检验"].copy()
    caption(doc, "表9  GM(1,1) 历史拟合值与相对误差")
    add_table(doc, fit, fontsize=10, col_widths=[2.4, 2.8, 2.8, 2.8, 3.0])
    body(doc, "计算得平均相对误差 MAPE=0.217%，后验差比值 C=0.0103（<0.35），小误差概率 P=1.000（>0.95），"
              "对照灰色模型精度等级标准，三项指标均达到一级（优），且 |a|=0.110<0.3，模型适合中短期预测，"
              "可用于外推。")
    h2(doc, "（三）2024—2028 年预测结果")
    fore = R["GM11预测"].copy()
    caption(doc, "表10  2024—2028 年全国数字经济综合发展指数预测值")
    add_table(doc, fore, fontsize=10.5, col_widths=[5.0, 5.0])
    add_figure(doc, "fig5_GM11预测.png", "图7  全国数字经济发展指数 GM(1,1) 拟合与预测（2024—2028）", width=14.0)
    body(doc, "预测结果（表10、图7）显示，全国数字经济综合发展指数将由 2023 年的 0.4551 上升至 2024 年的 "
              "0.5084，并持续增长到 2028 年的 0.7908，2023—2028 年年均增长率约 11.68%，整体延续强劲上升"
              "态势，与近年来数字经济与实体经济深度融合、新型基础设施加速布局的现实相吻合。需要说明的是，"
              "GM(1,1) 为趋势外推模型，预测值反映既有增长惯性下的基准情景，若出现重大技术变革或政策调整，"
              "应结合滚动数据及时修正。")

    # 九、问题四
    h1(doc, "九、问题四：分梯队差异化政策建议")
    body(doc, "综合评价、障碍诊断与趋势预测结果，本文从国家与三个梯队两个层面提出差异化建议。")
    h2(doc, "（一）第Ⅰ梯队：强化原始创新，打造数字经济增长极")
    body(doc, "北京、上海、广东、江苏、浙江应发挥资本、人才与数据要素集聚优势，重点突破基础软件、工业"
              "软件、核心算法与高端芯片等“卡脖子”环节，提升软件业务收入与高质量专利占比；同时面向高密度"
              "城市进一步加密 5G/算力等新型基础设施，探索数据要素市场化改革，形成可向全国复制的制度经验，"
              "当好数字经济“领头雁”。")
    h2(doc, "（二）第Ⅱ梯队：错位竞争，推动产业数字化攻坚")
    body(doc, "第Ⅱ梯队省份障碍集中于软件产业、电子商务和技术创新，应立足自身制造业基础与区位优势，"
              "推进“上云用数赋智”，提高关键工序数控化率和企业电子商务规模；承接第Ⅰ梯队产业、技术与人才"
              "外溢，培育特色数字产业集群，避免同质化竞争，力争整体向第Ⅰ梯队跃升。")
    h2(doc, "（三）第Ⅲ梯队：补齐底座，跨越区域数字鸿沟")
    body(doc, "西部和东北第Ⅲ梯队省份各维度普遍偏弱，首要任务是补齐全域网络覆盖、算力节点和数字人才"
              "短板。国家应通过中央转移支付、“东数西算”工程布局和对口协作，引导数据中心、呼叫服务、"
              "软件外包等产业向西部落地；优先发展远程教育、远程医疗和数字普惠金融，以数字普惠带动整体"
              "水平提升，防止在数字化进程中“掉队”。")
    h2(doc, "（四）国家层面：顶层设计与协同治理")
    body(doc, "国家层面应健全数字经济统计监测与区域协调机制，推动数据要素跨区域有序流动；针对障碍度"
              "最高的数字基础设施和数字产业化领域加大投入；预测期内全国数字经济仍将高速增长，应提前布局"
              "算力、能源与制度供给，统筹发展与安全，让数字经济红利更公平地惠及各地区。")

    # 十、模型评价与推广
    h1(doc, "十、模型的评价与推广")
    h2(doc, "（一）模型优点")
    for t in [
        "1. 方法组合科学、层次清晰。熵权法客观赋权与 TOPSIS 优劣排序相结合，避免了主观赋权偏差；系统聚类"
        "使梯队划分有数据依据；障碍度模型实现从“评价结果”到“成因诊断”的深化；GM(1,1) 契合小样本预测，"
        "形成完整闭环。",
        "2. 结果稳健、可复算。全部计算由程序自动完成、固定随机过程可复现，正文、图表与附录数据严格一致，"
        "并通过后验差比、小误差概率等多重检验保证可靠性。",
        "3. 可推广性强。所建“评价—诊断—预测”框架可直接迁移到城市数字经济、制造业高质量发展、生态环境"
        "质量等其他多指标综合评价场景，只需替换指标数据。"]:
        body(doc, t)
    h2(doc, "（二）模型不足")
    for t in [
        "1. 熵权法为客观赋权，虽避免了主观随意性，但无法体现指标的政策重要性差异，后续可与层次分析法、"
        "CRITIC 法组合赋权相互校验；",
        "2. 本文使用截面数据评价，尚未刻画省域数字经济的动态演进，未来可引入面板数据与马尔可夫链分析"
        "梯队转移概率；",
        "3. GM(1,1) 为单变量趋势模型，未纳入投资、政策等外生变量，长期预测精度会下降，可进一步采用 "
        "GM(1,N) 或灰色—马尔可夫组合模型改进。"]:
        body(doc, t)
    h2(doc, "（三）模型推广")
    body(doc, "将本文指标体系替换为相应领域的评价指标，即可把“熵权—TOPSIS—障碍度—灰色预测”方法链推广到"
              "区域营商环境评价、企业竞争力诊断、资源环境承载力预测等问题，具有较强的通用价值。")

    # 参考文献
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    h1(doc, "参考文献")
    refs = [
        "[1] 邓聚龙. 灰色控制系统[M]. 武汉: 华中工学院出版社, 1985.",
        "[2] 刘思峰, 党耀国, 方志耕, 等. 灰色系统理论及其应用[M]. 7版. 北京: 科学出版社, 2014.",
        "[3] Hwang C L, Yoon K. Multiple Attribute Decision Making: Methods and Applications "
        "[M]. Berlin: Springer-Verlag, 1981.",
        "[4] Shannon C E. A Mathematical Theory of Communication[J]. Bell System Technical "
        "Journal, 1948, 27(3): 379-423.",
        "[5] 郭显光. 改进的熵值法及其在经济效益评价中的应用[J]. 系统工程理论与实践, 1998, 18(12): 98-102.",
        "[6] Ward J H. Hierarchical Grouping to Optimize an Objective Function[J]. Journal of "
        "the American Statistical Association, 1963, 58(301): 236-244.",
        "[7] 赵涛, 张智, 梁上坤. 数字经济、创业活跃度与高质量发展——来自中国城市的经验证据[J]. 管理世界, "
        "2020, 36(10): 65-76.",
        "[8] 郭峰, 王靖一, 王芳, 等. 测度中国数字普惠金融发展: 指数编制与空间特征[J]. 经济学(季刊), "
        "2020, 19(4): 1401-1418.",
        "[9] 国家统计局. 中国统计年鉴2023[M]. 北京: 中国统计出版社, 2023.",
    ]
    for r in refs:
        p = doc.add_paragraph(); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.left_indent = Pt(24); p.paragraph_format.first_line_indent = Pt(-24)
        set_font(p.add_run(r), size=10.5)

    # 附录 A 原始数据
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("附录A  原始数据"), cn=CN_HEI, en="Arial", size=16, bold=True)
    raw = R["原始数据(省域截面)"].copy()
    raw.columns = ["省份"] + [f"X{i}" for i in range(1, 12)]
    caption(doc, "表A.1  31 省数字经济评价指标原始数据（2023 年）")
    add_table(doc, raw, fontsize=8.5)
    body(doc, "注：X1—X11 指标含义及单位见表2；X11 为负向指标。", indent=False)
    ts = R["原始数据(全国时序)"].copy()
    pagebreak(doc)
    caption(doc, "表A.2  全国数字经济综合发展指数（2014—2023）")
    add_table(doc, ts, fontsize=10, col_widths=[5.0, 7.0])

    # 附录 B 源代码
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("附录B  主要源程序代码（Python）"), cn=CN_HEI, en="Arial", size=16, bold=True)
    body(doc, "完整工程文件见随附 code/main.py，运行环境 Python 3.10 + numpy/pandas/scipy/matplotlib，"
              "在 code 目录执行 python main.py 即可一键复现本文全部表格与图形。核心代码如下：", indent=False)
    with open(os.path.join(BASE, "main.py"), "r", encoding="utf-8") as f:
        code_text = f.read()
    for line in code_text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Pt(0)
        set_font(p.add_run(line if line else " "), cn="Consolas", en="Consolas", size=8)

    # 参考文献/附录各节继承正文页码连续编号，去掉被复制过来的“起始页码=1”
    for s in doc.sections[4:]:
        el = s._sectPr.find(qn("w:pgNumType"))
        if el is not None:
            s._sectPr.remove(el)

    doc.save(OUT)
    print("报告已生成：", OUT)


if __name__ == "__main__":
    build()
